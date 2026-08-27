import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_ai_image_generation_excel(filename="AI画像生成管理表.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "画像生成プロンプト一覧"

    # Ensure gridlines are visible
    ws.views.sheetView[0].showGridLines = True

    # Title Banner (Optional nice header)
    ws.merge_cells("A1:G1")
    title_cell = ws["A1"]
    title_cell.value = "AI画像生成 管理シート"
    title_cell.font = Font(name="游ゴシック", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    # Subtitle / Note row
    ws.merge_cells("A2:G2")
    sub_cell = ws["A2"]
    sub_cell.value = "※ 「結果」欄には生成された画像を貼り付けるか、プロンプト結果を記入してください。"
    sub_cell.font = Font(name="游ゴシック", size=9, italic=True, color="595959")
    sub_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 20

    # Table Headers
    headers = ["No.", "テーマ", "ターゲット", "雰囲気", "構図", "生成プロンプト", "結果"]
    header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    header_font = Font(name="游ゴシック", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    header_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="medium", color="1F4E78"),
        bottom=Side(style="medium", color="1F4E78")
    )

    ws.row_dimensions[4].height = 28
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = header_border

    # Sample Data
    sample_data = [
        [
            1,
            "カフェの新商品紹介",
            "20代〜30代の女性",
            "温かみのあるナチュラルな雰囲気",
            "手元寄り・真上からの俯瞰ショット",
            "A cozy, warm-toned cafe scene featuring a seasonal latte on a wooden table, top-down view, natural lighting, soft focus background, aesthetic Instagram style, high quality, 8k photo",
            "【画像貼り付け欄】"
        ],
        [
            2,
            "フィットネスジムPR",
            "30代〜40代の社会人男女",
            "スタイリッシュ・アクティブ",
            "人物メイン・斜め前からのダイナミックな構図",
            "A stylish, dynamic photo of a young professional working out in a modern high-end gym, vibrant lighting, motivational vibe, professional photography",
            ""
        ],
        [
            3,
            "スキンケア商品の告知",
            "20代〜40代女性",
            "清潔感・透明感・ナチュラル",
            "商品中心・背景に水滴と光の反射",
            "Clean and minimal product photography of a serum bottle surrounded by subtle water ripples and gentle sunlight reflections, bright soft pastel tones, high resolution",
            ""
        ]
    ]

    # Fill sample rows
    start_row = 5
    num_blank_rows = 7  # Add blank template rows up to No. 10
    
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    fill_even = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_odd = PatternFill(start_color="F9FBFD", end_color="F9FBFD", fill_type="solid")

    font_regular = Font(name="游ゴシック", size=10, color="333333")
    font_no = Font(name="游ゴシック", size=10, bold=True, color="1F4E78")
    font_result_placeholder = Font(name="游ゴシック", size=9, color="A6A6A6", italic=True)

    current_row = start_row

    # Helper to apply row styles
    def apply_row_style(row_num, values, is_sample=True):
        ws.row_dimensions[row_num].height = 90  # Generous height for images
        row_fill = fill_even if (row_num % 2 == 0) else fill_odd

        for col_idx in range(1, 8):
            val = values[col_idx - 1] if col_idx - 1 < len(values) else ""
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.fill = row_fill
            cell.border = thin_border
            cell.font = font_regular

            # Column specific alignments & fonts
            if col_idx == 1:  # No.
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = font_no
            elif col_idx in [2, 3, 4, 5]:  # Themes, target, mood, composition
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif col_idx == 6:  # Prompt
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif col_idx == 7:  # Result (Image area)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if val == "【画像貼り付け欄】" or not val:
                    cell.font = font_result_placeholder

    # Apply sample data
    for item in sample_data:
        apply_row_style(current_row, item, is_sample=True)
        current_row += 1

    # Apply empty template rows up to 10
    for no in range(len(sample_data) + 1, 11):
        empty_item = [no, "", "", "", "", "", ""]
        apply_row_style(current_row, empty_item, is_sample=False)
        current_row += 1

    # Column Widths
    col_widths = {
        "A": 8,   # No.
        "B": 22,  # テーマ
        "C": 22,  # ターゲット
        "D": 22,  # 雰囲気
        "E": 25,  # 構図
        "F": 45,  # 生成プロンプト
        "G": 30   # 結果
    }

    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(filename)
    print(f"Successfully generated {filename}")

if __name__ == "__main__":
    create_ai_image_generation_excel()
